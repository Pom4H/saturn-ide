"""Independent, read-only verification of exported acceptance fixtures (openpyxl)."""
import warnings
from datetime import datetime
from pathlib import Path
import openpyxl

warnings.simplefilter("error")
root = Path(__file__).resolve().parent.parent / "artifacts"
book = openpyxl.load_workbook(root / "report-migration.xlsx")
assert book.sheetnames == ["Данные", "Рейтинг"]
sheet = book["Данные"]
assert sheet["B2"].value is None and sheet["B3"].value == 17.5
assert sheet["E2"].value is False and sheet["E3"].value is True
assert sheet["D2"].value == "=1+1" and sheet["D2"].data_type == "s"
assert isinstance(sheet["A2"].value, datetime)
assert sheet["A2"].value == datetime(2025, 12, 31, 23, 59, 59)
assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:E3"
assert sheet["B3"].number_format == "0.000"
assert sheet.column_dimensions["A"].width == 25
assert list(book["Рейтинг"].values) == [("Среднее (m³/h)",), (17.5,), (None,)]
assert sheet.oddFooter.left.text.startswith("Saturn")
browser = openpyxl.load_workbook(root / "report-browser.xlsx")
assert browser.sheetnames == ["Измерения", "Максимумы"]
assert browser["Измерения"]["B2"].value is None
assert browser["Измерения"].freeze_panes == "A2"
print("PASS independent XLSX read: sheets, types, exact dates/values/null, literal text, sort, formats, widths, filter, freeze and footer")
